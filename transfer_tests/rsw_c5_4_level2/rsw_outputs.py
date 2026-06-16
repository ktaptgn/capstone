"""Step 6 Excel, Markdown report, and presentation PNG generation."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rsw_c5_4_mini_sim import CONFIG_PATH, ROOT, load_config


POLICIES = ["H0", "H_TIME", "H1", "H2", "H3", "H4"]
COLORS = {
    "H0": "#9CA3AF",
    "H_TIME": "#6B7280",
    "H1": "#16A34A",
    "H2": "#0D9488",
    "H3": "#F59E0B",
    "H4": "#3B82F6",
}


def read_outputs(config: dict[str, Any]) -> dict[str, pd.DataFrame]:
    output_dir = ROOT / config["outputs"]["directory"]
    return {
        "summary": pd.read_csv(output_dir / config["outputs"]["policy_summary_csv"]),
        "event_log": pd.read_csv(output_dir / config["outputs"]["event_log_csv"], low_memory=False),
        "failure_log": pd.read_csv(output_dir / config["outputs"]["failure_log_csv"]),
        "mechanism": pd.read_csv(output_dir / config["outputs"]["mechanism_csv"]),
        "sensitivity": pd.read_csv(output_dir / config["outputs"]["sensitivity_csv"]),
    }


def build_rankings(summary: pd.DataFrame) -> pd.DataFrame:
    rankings = summary.copy()
    rankings["rank"] = rankings.groupby("regime")["TCO_mean"].rank(method="first")
    rankings["best_policy"] = rankings["rank"].eq(1)
    columns = [
        "regime",
        "rank",
        "best_policy",
        "policy",
        "TCO_mean",
        "TCO_std",
        "PMvisits_mean",
        "CM_mean",
        "failure_mean",
        "defect_mean",
        "fulfillment_mean",
    ]
    return rankings[columns].sort_values(["regime", "rank"])


def flatten_config(config_path: Path) -> pd.DataFrame:
    raw = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    rows: list[dict[str, Any]] = []

    def walk(value: Any, path: str) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                walk(child, f"{path}.{key}" if path else str(key))
        elif isinstance(value, list):
            rows.append({"config_path": path, "value": json.dumps(value, ensure_ascii=False)})
        else:
            rows.append({"config_path": path, "value": value})

    walk(raw, "")
    return pd.DataFrame(rows)


def interpretation_rows(summary: pd.DataFrame, sensitivity: pd.DataFrame) -> pd.DataFrame:
    rows = [
        {
            "topic": "scope",
            "interpretation": (
                "This RSW simulation is a synthetic Level 2 transfer mini-test. "
                "It is not a real factory validation."
            ),
        },
        {
            "topic": "main finding",
            "interpretation": (
                "H3 is best in heterogeneous_condition and high_stress; H4 is best in "
                "high_demand_high_stress. The expected mining ranking was not forced."
            ),
        },
        {
            "topic": "blind PM",
            "interpretation": (
                "H0 full-service calendar PM is the most expensive policy because frequent "
                "three-component visits dominate PM and downtime cost."
            ),
        },
        {
            "topic": "failure path",
            "interpretation": (
                "Only H_TIME produced failures in the 30-day base sweep; 10 of 12 occurred "
                "in the final campaign third."
            ),
        },
        {
            "topic": "component mechanism",
            "interpretation": (
                "State-aware and value/flow policies concentrate preventive maintenance on "
                "the electrode tip, while blind policies service all components equally."
            ),
        },
        {
            "topic": "slot sensitivity",
            "interpretation": (
                "Additional maintenance slots did not change the ranking. They increased blind "
                "full-service activity while H1-H4 were unchanged."
            ),
        },
    ]
    return pd.DataFrame(rows)


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    best_fill = PatternFill("solid", fgColor="C6E0B4")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 42)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        if sheet.title == "regime_rankings":
            headers = {cell.value: cell.column for cell in sheet[1]}
            best_column = headers.get("best_policy")
            if best_column:
                for row in range(2, sheet.max_row + 1):
                    if sheet.cell(row, best_column).value is True:
                        for cell in sheet[row]:
                            cell.fill = best_fill
    workbook.save(path)


def write_excel(config: dict[str, Any], data: dict[str, pd.DataFrame]) -> Path:
    output_dir = ROOT / config["outputs"]["directory"]
    path = output_dir / config["outputs"]["policy_summary_xlsx"]
    rankings = build_rankings(data["summary"])
    event_sample = data["event_log"].head(1000)
    config_df = flatten_config(CONFIG_PATH)
    interpretations = interpretation_rows(data["summary"], data["sensitivity"])
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        data["summary"].to_excel(writer, sheet_name="summary", index=False)
        rankings.to_excel(writer, sheet_name="regime_rankings", index=False)
        event_sample.to_excel(writer, sheet_name="event_log_sample", index=False)
        data["failure_log"].to_excel(writer, sheet_name="failure_log", index=False)
        data["mechanism"].to_excel(writer, sheet_name="mechanism", index=False)
        data["sensitivity"].to_excel(writer, sheet_name="sensitivity", index=False)
        config_df.to_excel(writer, sheet_name="config", index=False)
        interpretations.to_excel(writer, sheet_name="interpretation", index=False)
    style_workbook(path)
    return path


def _despine(ax) -> None:
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)


def write_plot(config: dict[str, Any], data: dict[str, pd.DataFrame]) -> Path:
    summary = data["summary"]
    mechanism = data["mechanism"]
    profiles = mechanism[mechanism["table"] == "profile"].copy()
    regimes = list(config["regimes"])

    fig, axes = plt.subplots(3, 1, figsize=(13, 14))
    width = 0.12
    x = range(len(regimes))
    for index, policy in enumerate(POLICIES):
        subset = summary[summary["policy"] == policy].set_index("regime").loc[regimes]
        positions = [value + (index - 2.5) * width for value in x]
        axes[0].bar(positions, subset["TCO_mean"], width, label=policy, color=COLORS[policy])
    axes[0].set_xticks(list(x), ["heterogeneous", "high stress", "high demand + stress"])
    axes[0].set_ylabel("TCO (synthetic CU)")
    axes[0].set_title("Policy TCO by regime (lower is better)")
    axes[0].legend(ncols=6, fontsize=8)
    _despine(axes[0])

    base = summary[summary["regime"] == "heterogeneous_condition"].set_index("policy").loc[POLICIES]
    positions = list(range(len(POLICIES)))
    axes[1].bar(
        [value - 0.18 for value in positions],
        base["PMvisits_mean"],
        0.36,
        label="PM visits",
        color="#2563EB",
    )
    axes[1].bar(
        [value + 0.18 for value in positions],
        base["CM_mean"],
        0.36,
        label="CM events",
        color="#DC2626",
    )
    axes[1].set_xticks(positions, POLICIES)
    axes[1].set_ylabel("Mean count per run")
    axes[1].set_title("PM visits and CM events (heterogeneous condition)")
    axes[1].legend()
    _despine(axes[1])

    profile = profiles[profiles["regime"] == "heterogeneous_condition"].set_index("policy").loc[POLICIES]
    bottom = pd.Series(0.0, index=POLICIES)
    component_colors = {"tip": "#F59E0B", "cooling": "#3B82F6", "actuator": "#8B5CF6"}
    for component in ("tip", "cooling", "actuator"):
        values = profile[f"pm_{component}_share"].astype(float)
        axes[2].bar(POLICIES, values, bottom=bottom, label=component, color=component_colors[component])
        bottom += values
    axes[2].set_ylim(0, 1.05)
    axes[2].set_ylabel("PM component share")
    axes[2].set_title("Preventive-maintenance component mix")
    axes[2].legend(ncols=3)
    _despine(axes[2])

    fig.suptitle(
        "RSW C5.4 Level 2 Synthetic Transfer Mini-Test\n"
        "Actual 10-seed x 30-day results; not a real factory validation",
        fontsize=14,
        fontweight="bold",
    )
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    path = ROOT / config["outputs"]["directory"] / config["outputs"]["policy_comparison_png"]
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return path


def fmt_rank_table(summary: pd.DataFrame) -> list[str]:
    lines = [
        "| regime | rank | policy | TCO | PM visits | CM | failure | defect | fulfillment |",
        "|---|---:|---|---:|---:|---:|---:|---:|---:|",
    ]
    ranked = build_rankings(summary)
    for _, row in ranked.iterrows():
        best = " **(best)**" if bool(row["best_policy"]) else ""
        lines.append(
            f"| {row['regime']} | {int(row['rank'])} | {row['policy']}{best} | "
            f"{row['TCO_mean']:.2f} | {row['PMvisits_mean']:.1f} | {row['CM_mean']:.2f} | "
            f"{row['failure_mean']:.2f} | {row['defect_mean']:.2f} | "
            f"{row['fulfillment_mean']:.3f} |"
        )
    return lines


def write_report(config: dict[str, Any], data: dict[str, pd.DataFrame]) -> Path:
    summary = data["summary"]
    sensitivity = data["sensitivity"]
    mechanism = data["mechanism"]
    profiles = mechanism[mechanism["table"] == "profile"]
    timing = mechanism[mechanism["table"] == "failure_timing"]
    htime_failures = int(timing[timing["policy"] == "H_TIME"]["total"].fillna(0).astype(float).sum())
    htime_late = int(timing[timing["policy"] == "H_TIME"]["late"].fillna(0).astype(float).sum())

    lines = [
        "# RSW C5.4 Level 2 Mini Transfer Simulation Report",
        "",
        "This RSW simulation is a synthetic Level 2 transfer mini-test. It is not a real factory "
        "validation. Its purpose is to reconstruct the C5.4 joint PM scheduling + dispatch "
        "structure in a manufacturing RSW setting.",
        "",
        "RSW 미니 시뮬레이션은 본 프로젝트의 메인 실험이 아니라 제조업 전이 가능성을 보여주는 "
        "보조 실험이다. C5.4의 truck/component/route/PM bay 구조를 RSW의 welding "
        "gun/component/job family/maintenance slot 구조로 재구성했다. 실제 현장 적용을 "
        "위해서는 용접 로그, 전극 dressing 이력, 품질 검사 데이터로 파라미터 보정이 필요하다.",
        "",
        "## 1. Purpose",
        "",
        "Demonstrate that the C5.4 joint PM scheduling and dispatch decision structure can be "
        "reconstructed as a reduced synthetic RSW mini-test.",
        "",
        "## 2. Why this is a transfer mini-test, not the main project",
        "",
        "The main project remains the mining C5.4 experiment. RSW parameters use synthetic CU and "
        "are not calibrated to a factory, vehicle program, or real maintenance standard.",
        "",
        "## 3. C5.4-to-RSW mapping",
        "",
        "| C5.4 mining | RSW mini transfer |",
        "|---|---|",
        "| Truck | Welding gun/cell |",
        "| Tire / engine / brake HI | Tip / cooling / actuator HI |",
        "| Route A/B/C | Job family A/B/C |",
        "| PM bay | Maintenance/dressing slot |",
        "| Dispatch | Production assignment |",
        "| TCO | PM + CM + downtime + degradation + unmet + defect |",
        "",
        "## 4. Simulation environment",
        "",
        "- 10 guns, 2 maintenance slots, 30 days, 24 hourly steps/day, 10 seeds.",
        "- Daily demand: 210 welds split equally across job families A/B/C.",
        "- Regimes: heterogeneous condition, high stress, high demand + high stress.",
        "",
        "## 5. Reliability surface",
        "",
        "Each gun has tip, cooling, and actuator HI; gun HI is the component minimum. Gamma "
        "frailty, noisy observed HI, partial PM/CM restoration, defect events, and threshold-power "
        "failure hazard are enabled.",
        "",
        "## 6. Joint PM scheduling + production assignment",
        "",
        "Each policy returns PM referrals and ranked job families in one hourly decision. PM "
        "referrals are constrained by free gun-level maintenance slots.",
        "",
        "## 7. Policy definitions: H0-H4 + H_TIME",
        "",
        "- H0: calendar full service + blind round-robin.",
        "- H_TIME: operating-hours full service + blind round-robin.",
        "- H1: condition CBM + health routing.",
        "- H2: risk-priority PM + risk-aware routing.",
        "- H3: cost-value PM + value/risk routing.",
        "- H4: flow-backpressure PM + capacity routing.",
        "",
        "## 8. Regime results",
        "",
        "Lower TCO is better. Actual results are retained even when they differ from the expected "
        "C5.4 mining ranking.",
        "",
        *fmt_rank_table(summary),
        "",
        "## 9. Sensitivity results",
        "",
        "- Frailty-CV sweep: the lowest-TCO policy changes across H1, H4, H3, and H1.",
        "- Best state-aware H1/H2 remains roughly 77-80% cheaper than best blind H0/H_TIME.",
        "- H1 CBM threshold perturbation changes TCO modestly, but the best scale differs by regime.",
        "- Extra maintenance slots do not change ranking; they increase blind full-service activity.",
        "",
        "## 10. Mechanism analysis",
        "",
        f"- H_TIME generated {htime_failures} failures; {htime_late} occurred in the late campaign third.",
        "- H0 performs 100 full-service visits/run and has the highest TCO despite near-zero failures.",
        "- H1-H4 concentrate PM on the electrode tip, while blind policies split PM equally.",
        "- Job-family completion shares are equal because demand is explicitly split A/B/C.",
        "",
        "## 11. Key findings",
        "",
        "- H3 is best under heterogeneous condition and high stress.",
        "- H4 is best under high demand + high stress.",
        "- Blind periodic PM is substantially more expensive, primarily from full-service PM and downtime.",
        "- The mini-test does not reproduce the C5.4 mining H1/H2-best ranking; this difference is an honest result.",
        "",
        "## 12. Limitations",
        "",
        "- Synthetic parameters and synthetic CU; no factory calibration.",
        "- Reduced 30-day, 10-seed horizon.",
        "- Equal job-family demand forces equal completed route shares, limiting dispatch-mix interpretation.",
        "- Failures are sparse, so frailty/failure correlation is undefined for most policies.",
        "- No PPO/RL training and no operator dashboard.",
        "",
        "## 13. Final presentation usage",
        "",
        "Use this mini-test as supporting evidence that the joint decision architecture transfers to "
        "manufacturing. Do not present it as proof of real-factory performance.",
    ]
    path = ROOT / config["outputs"]["directory"] / config["outputs"]["report_markdown"]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def main() -> int:
    parser = argparse.ArgumentParser(description="Build RSW Level 2 Step 6 deliverables.")
    parser.parse_args()
    config = load_config()
    data = read_outputs(config)
    excel = write_excel(config, data)
    plot = write_plot(config, data)
    report = write_report(config, data)
    print(f"Wrote {excel}")
    print(f"Wrote {plot}")
    print(f"Wrote {report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
